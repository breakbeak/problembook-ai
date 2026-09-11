# -*- coding: utf-8 -*-
import os, secrets, json, base64, tempfile, pathlib, sqlite3, hashlib, hmac, requests
from flask import Flask, request, jsonify, render_template, session, redirect, url_for

app = Flask(__name__)
app.secret_key = os.environ.get("ADMIN_SESSION_SECRET", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

MODEL = os.environ.get("OPENROUTER_MODEL", "openrouter/free")
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
DATA_DIR = pathlib.Path("data")
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_FILE = DATA_DIR / "app.db"


def db():
    c = sqlite3.connect(DB_FILE, timeout=10)
    c.row_factory = sqlite3.Row
    return c


def hash_pw(pw):
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 120000).hex()
    return salt + "$" + digest


def check_pw(pw, stored):
    try:
        salt, digest = stored.split("$", 1)
        actual = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 120000).hex()
        return hmac.compare_digest(actual, digest)
    except Exception:
        return False


def init_db():
    c = db()
    c.execute("""CREATE TABLE IF NOT EXISTS users(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        is_admin INTEGER DEFAULT 0)""")
    c.execute("""CREATE TABLE IF NOT EXISTS settings(
        name TEXT PRIMARY KEY,
        value TEXT NOT NULL)""")
    if c.execute("SELECT COUNT(*) FROM users WHERE is_admin=1").fetchone()[0] == 0:
        pw = os.environ.get("ADMIN_PASSWORD", "change-me")
        c.execute("INSERT OR IGNORE INTO users(username,password_hash,is_admin) VALUES(?,?,1)",
                  ("admin", hash_pw(pw)))
    c.commit(); c.close()


init_db()


def get_api_key():
    env = os.environ.get("OPENROUTER_API_KEY", "").strip()
    if env:
        return env
    c = db()
    row = c.execute("SELECT value FROM settings WHERE name='openrouter_api_key'").fetchone()
    if not row:
        # Backward compatibility with the previous version.
        row = c.execute("SELECT value FROM settings WHERE name='openai_api_key'").fetchone()
    c.close()
    return row["value"].strip() if row else ""


def admin_required():
    return session.get("admin") is True


def openrouter_request(key, messages, max_tokens=3000):
    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "https://problembook-ai.onrender.com",
        "X-Title": "AI Problembook",
    }
    payload = {"model": MODEL, "messages": messages, "max_tokens": max_tokens}
    r = requests.post(OPENROUTER_URL, headers=headers, json=payload, timeout=120)
    try:
        data = r.json()
    except Exception:
        raise RuntimeError(f"OpenRouter 응답 오류 ({r.status_code}): {r.text[:500]}")
    if r.status_code >= 400:
        err = data.get("error", {})
        raise RuntimeError(err.get("message") or f"OpenRouter 오류 ({r.status_code})")
    try:
        return data["choices"][0]["message"]["content"]
    except Exception:
        raise RuntimeError("AI 응답 형식이 올바르지 않습니다.")


def extract_text(path, ext):
    if ext in {".txt", ".md", ".csv"}:
        return pathlib.Path(path).read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:
            raise RuntimeError("PDF 텍스트 읽기 실패: " + str(e))
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        out = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    out.append(shape.text)
        return "\n".join(out)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    return ""


def clean_json_text(raw):
    raw = raw.strip()
    if raw.startswith("```"):
        raw = raw.strip("`").strip()
        if raw.lower().startswith("json"):
            raw = raw[4:].strip()
    try:
        return json.loads(raw)
    except Exception:
        start, end = raw.find("{"), raw.rfind("}")
        if start >= 0 and end > start:
            return json.loads(raw[start:end + 1])
        raise RuntimeError("AI가 JSON 형식으로 답하지 않았습니다.")


@app.get("/health")
def health():
    return jsonify(ok=True)


@app.get("/")
def index():
    if not session.get("user") and not session.get("admin"):
        return redirect(url_for("user_login"))
    return render_template("user.html")


@app.get("/login")
def user_login():
    if session.get("user") or session.get("admin"):
        return redirect(url_for("index"))
    return render_template("user_login.html")


@app.post("/login")
def user_login_post():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    c = db(); u = c.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone(); c.close()
    if u and check_pw(password, u["password_hash"]):
        session["user"] = username
        return redirect(url_for("index"))
    return render_template("user_login.html", error="아이디 또는 비밀번호가 틀렸어요.")


@app.post("/logout")
def user_logout():
    session.clear()
    return redirect(url_for("user_login"))


@app.get("/admin")
def admin():
    if not admin_required():
        return render_template("login.html")
    return render_template("admin.html", has_key=bool(get_api_key()), model=MODEL)


@app.post("/admin/login")
def admin_login():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    c = db(); u = c.execute("SELECT * FROM users WHERE username=? AND is_admin=1", (username,)).fetchone(); c.close()
    if u and check_pw(password, u["password_hash"]):
        session["admin"] = True
        session["admin_user"] = username
        return redirect(url_for("admin"))
    return render_template("login.html", error="아이디 또는 비밀번호가 틀렸어요.")


@app.post("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin"))


@app.post("/admin/save-key")
def save_key():
    if not admin_required():
        return jsonify(ok=False, error="관리자 로그인이 필요합니다."), 401
    key = request.form.get("api_key", "").strip()
    if not key:
        return jsonify(ok=False, error="API 키를 입력해 주세요."), 400
    try:
        c = db()
        c.execute("""INSERT INTO settings(name,value) VALUES('openrouter_api_key',?)
                     ON CONFLICT(name) DO UPDATE SET value=excluded.value""", (key,))
        c.commit(); c.close()
        return jsonify(ok=True, message="OpenRouter API 키가 저장됐습니다!")
    except Exception as e:
        return jsonify(ok=False, error="저장 오류: " + str(e)), 500


@app.post("/admin/test")
def test_key():
    if not admin_required():
        return jsonify(ok=False, error="관리자 로그인이 필요합니다."), 401
    key = get_api_key()
    if not key:
        return jsonify(ok=False, error="먼저 API 키를 저장해 주세요."), 400
    try:
        raw = openrouter_request(key, [{"role": "user", "content": "Reply with exactly: OK"}], max_tokens=10)
        return jsonify(ok=True, result=raw.strip())
    except Exception as e:
        return jsonify(ok=False, error=str(e)), 400


@app.post("/admin/create-user")
def create_user():
    if not admin_required():
        return jsonify(ok=False, error="관리자 로그인이 필요합니다."), 401
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    if len(username) < 3 or len(password) < 6:
        return jsonify(ok=False, error="아이디는 3자 이상, 비밀번호는 6자 이상이어야 합니다."), 400
    try:
        c = db()
        c.execute("INSERT INTO users(username,password_hash,is_admin) VALUES(?,?,0)",
                  (username, hash_pw(password)))
        c.commit(); c.close()
        return jsonify(ok=True, message="계정이 저장됐습니다!")
    except sqlite3.IntegrityError:
        return jsonify(ok=False, error="이미 존재하는 아이디입니다."), 400


@app.post("/api/analyze")
def analyze():
    if not session.get("user") and not session.get("admin"):
        return jsonify(error="로그인이 필요합니다."), 401
    key = get_api_key()
    if not key:
        return jsonify(error="관리자가 OpenRouter API 키를 설정하지 않았습니다."), 503
    files = request.files.getlist("files")
    if not files:
        return jsonify(error="파일을 하나 이상 올려주세요."), 400

    allowed = {".jpg", ".jpeg", ".png", ".webp", ".pdf", ".txt", ".md", ".csv", ".docx", ".pptx", ".xlsx"}
    temp_paths = []
    try:
        content = []
        text_parts = []
        for f in files:
            if not f.filename:
                continue
            ext = pathlib.Path(f.filename).suffix.lower()
            if ext not in allowed:
                return jsonify(error=f"지원하지 않는 파일 형식: {ext}"), 400
            fd, path = tempfile.mkstemp(suffix=ext); os.close(fd)
            f.save(path); temp_paths.append(path)
            if ext in {".jpg", ".jpeg", ".png", ".webp"}:
                mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}[ext]
                b64 = base64.b64encode(pathlib.Path(path).read_bytes()).decode()
                content.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}})
            else:
                text = extract_text(path, ext)
                if text.strip():
                    text_parts.append(f"[파일: {f.filename}]\n{text[:120000]}")

        prompt = """너는 중학생용 문제집 AI 튜터다.
업로드된 문제를 읽고 반드시 JSON 객체 하나만 반환해라. 마크다운 코드블록을 사용하지 마라.
형식:
{
 "summary":"문제 요약",
 "problems":[
 {"number":"1","problem":"인식한 문제","answer":"정답","method1":"쉬운 풀이",
 "method2":"다른 풀이","hint1":"첫 힌트","hint2":"두 번째 힌트",
 "concepts":["개념1"],"similar":"비슷한 문제"}
 ]
}
문제가 여러 개면 순서대로 모두 넣어라. 문제에 없는 조건은 만들지 마라."""
        if text_parts:
            content.append({"type": "text", "text": "\n\n".join(text_parts)})
        if not content:
            return jsonify(error="읽을 수 있는 파일 내용이 없습니다."), 400

        messages = [{"role": "user", "content": [{"type": "text", "text": prompt}] + content}]
        raw = openrouter_request(key, messages, max_tokens=6000)
        data = clean_json_text(raw)
        return jsonify(data)
    except Exception as e:
        return jsonify(error=str(e)), 500
    finally:
        for path in temp_paths:
            try:
                os.remove(path)
            except Exception:
                pass


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8080")))
